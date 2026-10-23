"""End-to-end smoke test (no pytest needed):  python tests/smoke_test.py
Runs the real app against a throw-away database and walks the whole flow."""
import os
import re
import sys
import tempfile
import time
from datetime import date, timedelta

tmp = tempfile.mkdtemp(prefix="itduty-test-")
os.environ["DATA_DIR"] = tmp
os.environ["COOKIE_SECURE"] = "0"
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app import main  # noqa: E402  (must come after the environment is set)

app = main.app
failures = []


def check(label, ok):
    print(("PASS  " if ok else "FAIL  ") + label)
    if not ok:
        failures.append(label)


def token(client, path="/admin/login"):
    html = client.get(path).get_data(as_text=True)
    m = re.search(r'name="_csrf" value="([^"]+)"', html)
    return m.group(1) if m else ""


def login(client, user, pw):
    return client.post("/admin/login", data={"_csrf": token(client), "username": user, "password": pw})


OWNER_PW = "Correct-Horse-7"
runner = app.test_cli_runner()

check("phone format: white space tidied, nothing else touched", main.fmt_phone("  511   22-33 44 ") == "511 22-33 44" and main.fmt_phone("511223344") == "511223344")
check("phone format: international numbers and empty values", main.fmt_phone("+1 555 0100") == "+1 555 0100" and main.fmt_phone("12-34") == "12-34" and main.fmt_phone(None) == "")
# --- Stage 1: Georgian holidays (Labour Code art. 20 + the Orthodox Easter days), computed without internet
from app import holidays as _hol
check("easter: Orthodox Easter 2026 is 12 April", _hol.orthodox_easter(2026) == date(2026, 4, 12))
check("easter: 2027 is 2 May, 2025 is 20 April, 2028 is 16 April", _hol.orthodox_easter(2027) == date(2027, 5, 2) and _hol.orthodox_easter(2025) == date(2025, 4, 20) and _hol.orthodox_easter(2028) == date(2028, 4, 16))
_h26 = _hol.auto_holidays(2026)
check("holidays: the 13 fixed dates are there (Giorgoba on 23 November 2026)", _h26[date(2026, 11, 23)] == "St. George's Day" and date(2026, 1, 1) in _h26 and date(2026, 1, 2) in _h26 and date(2026, 1, 7) in _h26 and date(2026, 1, 19) in _h26 and date(2026, 3, 3) in _h26 and date(2026, 3, 8) in _h26 and date(2026, 4, 9) in _h26 and date(2026, 5, 9) in _h26 and date(2026, 5, 12) in _h26 and date(2026, 5, 26) in _h26 and date(2026, 8, 28) in _h26 and date(2026, 10, 14) in _h26)
check("holidays: the four Easter days move with Easter (2026: 10, 11, 12 and 13 April)", all(date(2026, 4, d) in _h26 for d in (10, 11, 12, 13)) and _h26[date(2026, 4, 10)] == "Good Friday" and _h26[date(2026, 4, 13)] == "Easter Monday")
check("holidays: Easter 2027 days are 30 April, 1, 2 and 3 May", all(d in _hol.auto_holidays(2027) for d in (date(2027, 4, 30), date(2027, 5, 1), date(2027, 5, 2), date(2027, 5, 3))))
check("holidays: an ordinary day is not a holiday", date(2026, 11, 24) not in _h26)
import sqlite3 as _sq3
from app import db as _db
_old = _sq3.connect(":memory:")
_old.row_factory = _sq3.Row
_old.executescript("""
CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT);
CREATE TABLE staff (id INTEGER PRIMARY KEY, name TEXT NOT NULL, email TEXT NOT NULL DEFAULT '', phone TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE duty (day TEXT PRIMARY KEY, staff_id INTEGER NOT NULL REFERENCES staff(id));
INSERT INTO staff(id, name) VALUES (1, 'a'), (2, 'b');
INSERT INTO duty(day, staff_id) VALUES ('2026-10-21', 1), ('2026-10-22', 2);
""")
_db.migrate(_old)
check("migration: staff gets a branch column, empty for everybody", [r["branch"] for r in _old.execute("SELECT branch FROM staff ORDER BY id")] == ["", ""])
check("migration: old duty rows survive in the 'branch not set' slot", [(r["day"], r["branch"], r["staff_id"]) for r in _old.execute("SELECT * FROM duty ORDER BY day")] == [("2026-10-21", "", 1), ("2026-10-22", "", 2)])
check("migration: the same day can now hold two people (site_a and site_b)", _old.execute("INSERT INTO duty(day, branch, staff_id) VALUES ('2026-10-21', 'site_a', 2)").rowcount == 1)
_db.migrate(_old)
check("migration: running it twice changes nothing", _old.execute("SELECT COUNT(*) FROM duty").fetchone()[0] == 3)

# --- Task 2: short labels for calendar cells
_names = ["John Lowell", "Nina Baker", "Anna Kemp"]
check("short: unique first name -> first name + surname initial", main.cell_label("John Lowell", _names) == "John L.")
_dup = ["John Lowell", "John Loring", "John Baker", "Nina Baker"]
check("short: same first name and initial -> longer surname part", main.cell_label("John Lowell", _dup) == "John Low." and main.cell_label("John Loring", _dup) == "John Lor.")
check("short: same first name, different initial stays short", main.cell_label("John Baker", _dup) == "John B.")
check("short: single-word name unchanged", main.cell_label("Admin", _names) == "Admin")
check("short: identical full names cannot be told apart, full surname shown", main.cell_label("Nina Baker", ["Nina Baker", "Nina Baker"]) == "Nina Baker")
# --- Task 1: week groups for the public page (pure logic, no database needed)
from datetime import date as _d
tw, nw = main.week_groups(_d(2026, 10, 21), {}, {})           # a Wednesday
check("weeks: Wednesday -> this week is Wed..Sun (5 days)", len(tw) == 5 and tw[0]["date"] == _d(2026, 10, 21) and tw[-1]["date"] == _d(2026, 10, 25))
check("weeks: next week is the full Mon..Sun", len(nw) == 7 and nw[0]["date"] == _d(2026, 10, 26) and nw[-1]["date"] == _d(2026, 11, 1))
check("weeks: today and tomorrow flags", tw[0]["is_today"] and tw[1]["is_tomorrow"] and not tw[2]["is_tomorrow"])
check("weeks: Saturday and Sunday are weekend", [x["weekend"] for x in tw] == [False, False, False, True, True])
sw, sn = main.week_groups(_d(2026, 10, 25), {}, {})           # a Sunday
check("weeks: on a Sunday this week is only today", len(sw) == 1 and sw[0]["is_today"] and sn[0]["is_tomorrow"])
yw, yn = main.week_groups(_d(2026, 12, 30), {}, {})           # year rollover
_hw, _ = main.week_groups(_d(2026, 10, 21), {("2026-10-22", "site_a"): [{"name": "e"}, {"name": "e2"}], ("2026-10-22", "site_b"): [{"name": "w"}]}, {_d(2026, 10, 23): "test"})
check("weeks: site_a and site_b of a day, holiday name", [p["name"] for p in _hw[1]["site_a"]] == ["e", "e2"] and _hw[1]["site_b"][0]["name"] == "w" and _hw[0]["site_a"] == [] and _hw[2]["holiday"] == "test" and _hw[0]["holiday"] is None)
check("weeks: year rollover is handled", yw[-1]["date"] == _d(2027, 1, 3) and yn[0]["date"] == _d(2027, 1, 4))

# --- first run helper: admin/admin that MUST be changed at first login
res = runner.invoke(args=["bootstrap-admin"])
check("CLI bootstrap-admin creates the initial owner", res.exit_code == 0 and "'admin'" in res.output)
boot = app.test_client()
check("admin/admin logs in", login(boot, "admin", "admin").status_code == 302)
check("...but is forced to change the password first", "settings" in boot.get("/admin/schedule").headers["Location"])
check("a too-short password is refused (min 7)", "at least 7" in boot.post("/admin/settings", data={"_csrf": token(boot, "/admin/settings"), "current": "admin", "new": "admin1", "again": "admin1"}).get_data(as_text=True))
boot.post("/admin/settings", data={"_csrf": token(boot, "/admin/settings"), "current": "admin", "new": "My-First-Real-Pass-1", "again": "My-First-Real-Pass-1"})
check("after the change the old password 'admin' is dead", login(app.test_client(), "admin", "admin").status_code == 401)
check("bootstrap-admin does nothing once accounts exist", "nothing to do" in runner.invoke(args=["bootstrap-admin"]).output)

# --- first run: create the owner through the CLI (this is how it is done in the container)
res = runner.invoke(args=["create-owner"], input=f"owner\nTest Owner\n{OWNER_PW}\n{OWNER_PW}\n")
check("CLI create-owner", res.exit_code == 0 and "created" in res.output)
res = runner.invoke(args=["seed-demo"])
check("CLI seed-demo", res.exit_code == 0)

owner = app.test_client()
anon = app.test_client()

# --- public side
r = anon.get("/")
body = r.get_data(as_text=True)
check("public page 200 + title", r.status_code == 200 and "IT Duty" in body)
check("public page shows someone on duty", "On duty now" in body)
check("public page: the on-duty pill comes before the name, no old label",
      'class="eyebrow"' not in body and body.index("On duty now") < body.index('class="duty-name"'))
check("public page: new title in the header and both week groups", "IT Duty" in body and "This week" in body and "Next week" in body)
check("public page: weekend rows are marked", 'wk-row we' in body)
check("public page has no Call button and no initials avatar", "Call" not in body and 'class="avatar"' not in body)
check("security headers set", "default-src 'self'" in r.headers.get("Content-Security-Policy", ""))
check("healthz ok", anon.get("/healthz").status_code == 200)
check("public page: the admin link is a plain link named admin panel (no popup script)", ">Admin panel</a>" in body and "data-admin-link" not in body and 'id="toast"' not in body)
check("the old /denied page no longer exists (404)", anon.get("/denied").status_code == 404)
check("api/today json: a list of people per branch, no e-mail", isinstance(anon.get("/api/today").get_json()["on_duty"]["site_a"], list) and "email" not in anon.get("/api/today").get_data(as_text=True))
check("admin requires login (redirect)", anon.get("/admin/schedule").status_code == 302)
check("POST without CSRF rejected", anon.post("/admin/login", data={"username": "owner", "password": OWNER_PW}).status_code == 400)

# --- login
check("wrong password -> 401", login(owner, "owner", "wrong-password-123").status_code == 401)
check("right password -> redirect", login(owner, "owner", OWNER_PW).status_code == 302)
check("schedule page renders", owner.get("/admin/schedule").status_code == 200)
check("week view renders", owner.get("/admin/schedule?view=week").status_code == 200)
for page in ("/admin/staff", "/admin/users", "/admin/audit", "/admin/settings"):
    check(f"owner can open {page}", owner.get(page).status_code == 200)

# --- schedule + staff
t = owner.get("/admin/staff")
tok = token(owner, "/admin/staff")
owner.post("/admin/staff/save", data={"_csrf": tok, "name": "New Person", "branch": "site_a", "email": "new@example.com", "phone": "+1 555 0100"})
check("staff added", "New Person" in owner.get("/admin/staff").get_data(as_text=True))
with app.app_context():
    sid = main.q1("SELECT id FROM staff WHERE name = ?", "New Person")["id"]
far = (date.today() + timedelta(days=3)).isoformat()
owner.post("/admin/schedule/assign", data={"_csrf": token(owner, "/admin/schedule"), "day": far, "staff_id": sid})
with app.app_context():
    check("duty assigned", sid in [r["staff_id"] for r in main.q("SELECT staff_id FROM duty WHERE day = ? AND branch = 'site_a'", far)])
owner.post("/admin/schedule/range", data={"_csrf": token(owner, "/admin/schedule"), "staff_id": sid,
                                          "from": far, "to": (date.today() + timedelta(days=9)).isoformat()})
with app.app_context():
    check("range: every day gets the person, weekends included (7 days)", main.q1("SELECT COUNT(*) AS n FROM duty WHERE staff_id = ? AND day BETWEEN ? AND ?", sid, far, (date.today() + timedelta(days=9)).isoformat())["n"] == 7)
owner.post("/admin/schedule/copy-week", data={"_csrf": token(owner, "/admin/schedule")})
check("bad staff id rejected", owner.post("/admin/schedule/assign", data={"_csrf": token(owner, "/admin/schedule"), "day": far, "staff_id": "9999"}).status_code == 400)

# --- Task 3: swap
owner.post("/admin/staff/save", data={"_csrf": token(owner, "/admin/staff"), "name": "Second Person", "branch": "site_a", "email": "", "phone": ""})
with app.app_context():
    sid2 = main.q1("SELECT id FROM staff WHERE name = ?", "Second Person")["id"]
da, db_, dc = [(date.today() + timedelta(days=n)).isoformat() for n in (40, 41, 42)]
for dd, ss in ((da, sid), (db_, sid2)):
    owner.post("/admin/schedule/assign", data={"_csrf": token(owner, "/admin/schedule"), "day": dd, "staff_id": ss})
def swap(a, b):
    return owner.post("/admin/schedule/swap", data={"_csrf": token(owner, "/admin/schedule"), "a": a, "b": b, "branch": "site_a"})
def duty_of(dd, branch="site_a"):
    with app.app_context():
        r = main.q1("SELECT staff_id FROM duty WHERE day = ? AND branch = ?", dd, branch)
        return r["staff_id"] if r else None
def audit_count(action):
    with app.app_context():
        return main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = ?", action)["n"]
before = audit_count("duty.swap")
swap(da, db_)
check("swap: two assigned days trade places", duty_of(da) == sid2 and duty_of(db_) == sid)
check("swap: exactly one audit entry", audit_count("duty.swap") == before + 1)
swap(da, dc)
check("swap: into an empty day moves the person", duty_of(da) is None and duty_of(dc) == sid2)
check("swap: same day is refused (400)", swap(db_, db_).status_code == 400 and duty_of(db_) == sid)
check("swap: two empty days refused (400)", swap(da, (date.today() + timedelta(days=50)).isoformat()).status_code == 400)
check("swap: bad date refused (400)", swap("nonsense", da).status_code == 400)
check("swap: bad branch refused (400)", owner.post("/admin/schedule/swap", data={"_csrf": token(owner, "/admin/schedule"), "a": da, "b": dc, "branch": "north"}).status_code == 400)
check("swap: needs login", anon.post("/admin/schedule/swap", data={"a": da, "b": db_}).status_code in (302, 400, 403))
with app.app_context():
    last = main.q1("SELECT detail FROM audit WHERE action = 'duty.swap' ORDER BY id DESC LIMIT 1")["detail"]
check("swap: audit text names both days", dc in last and da in last)
page = owner.get("/admin/schedule").get_data(as_text=True)
check("schedule page: people palette lists staff with counts", 'id="palette"' in page and "Second Person" in page and 'data-staff-id="%d"' % sid2 in page)

# --- bulk select: assign or unassign several days at once
bd = [(date.today() + timedelta(days=n)).isoformat() for n in (60, 61, 62, 64)]
def bulk(days, staff_id):
    return owner.post("/admin/schedule/bulk", data={"_csrf": token(owner, "/admin/schedule"), "days": ",".join(days), "staff_id": staff_id, "branch": "site_a"})
before = audit_count("duty.bulk")
bulk(bd, str(sid))
check("bulk: one person assigned to all selected days", all(duty_of(x) == sid for x in bd))
check("bulk: exactly one audit entry", audit_count("duty.bulk") == before + 1)
bulk(bd[:2], "")
check("bulk: unassign clears the selected days only", duty_of(bd[0]) is None and duty_of(bd[1]) is None and duty_of(bd[2]) == sid and duty_of(bd[3]) == sid)
check("bulk: clearing already-empty days is harmless", bulk(bd[:2], "").status_code in (302, 200))
check("bulk: bad date refused (400), nothing changed", bulk([bd[2], "nonsense"], "").status_code == 400 and duty_of(bd[2]) == sid)
check("bulk: empty selection refused (400)", bulk([], str(sid)).status_code == 400)
check("bulk: unknown staff refused (400)", bulk(bd, "9999").status_code == 400)
many = [(date.today() + timedelta(days=100 + n)).isoformat() for n in range(120)]
check("bulk: more than 100 days refused (400)", bulk(many, str(sid)).status_code == 400 and duty_of(many[0]) is None)
check("bulk: duplicate days counted once", bulk([bd[0], bd[0], bd[0]], str(sid)).status_code in (302, 200) and duty_of(bd[0]) == sid)
check("schedule page: multi-select controls present", 'id="multi-toggle"' in owner.get("/admin/schedule").get_data(as_text=True) and 'id="frm-bulk"' in owner.get("/admin/schedule").get_data(as_text=True))

# --- Task 4: Excel paste import + xlsx export
from app import xlsx_io
import io as _io
_staff = [{"id": 1, "name": "John Lowell", "branch": "site_a"}, {"id": 2, "name": "Nina Baker", "branch": "site_a"}, {"id": 3, "name": "Nina Baker-Kemp", "branch": "site_a"}]
r2 = xlsx_io.parse_paste("2026-11-02\tJohn Lowell\n03.11.2026\tNina Baker\n04/11/2026\tJohn Lowell", _staff)
check("paste: ISO, DD.MM.YYYY and DD/MM/YYYY all parse", [(d.isoformat(), i) for d, i in r2["ok"]] == [("2026-11-02", 1), ("2026-11-03", 2), ("2026-11-04", 1)] and not r2["bad"])
r3 = xlsx_io.parse_paste("\n\n  2026-11-05  \t  John   Lowell  \r\n\r\n", _staff)
check("paste: stray spaces, blank lines and CRLF are tolerated", [(d.isoformat(), i) for d, i in r3["ok"]] == [("2026-11-05", 1)] and not r3["bad"])
r4 = xlsx_io.parse_paste("John Lowell\t2026-11-06\nLowell John\t07.11.2026\nBaker Nina\t08.11.2026", _staff)
check("paste: name before date and surname-first order both work", [(d.isoformat(), i) for d, i in r4["ok"]]== [("2026-11-06", 1), ("2026-11-07", 1), ("2026-11-08", 2)])
r5 = xlsx_io.parse_paste("2026-11-09\tUnknown Person\n2026-02-31\tJohn Lowell\nno date at all\n2026-11-10\n", _staff)
check("paste: unknown names, impossible dates, lines without date or name are skipped with a reason", not r5["ok"] and len(r5["bad"]) == 4 and all(b[0] > 0 and b[2] for b in r5["bad"]))
r6 = xlsx_io.parse_paste("2026-11-11\tNina Baker\n2026-11-12\tNina Baker-Kemp", _staff)
check("paste: a name that is a prefix of another matches exactly", [i for _, i in r6["ok"]] == [2, 3])
r7 = xlsx_io.parse_paste("2026-11-13\tJohn Lowell\n2026-11-13\tNina Baker", _staff)
check("paste: two people on the same date and branch are both accepted", [i for _, i in r7["ok"]] == [1, 2] and not r7["bad"])
r7b = xlsx_io.parse_paste("2026-11-13\tJohn Lowell\n2026-11-13\tJohn Lowell", _staff)
check("paste: the same person twice on one date is reported once", [i for _, i in r7b["ok"]] == [1] and len(r7b["bad"]) == 1)
_mixed = _staff[:1] + [{"id": 9, "name": "Site B Person", "branch": "site_b"}, {"id": 8, "name": "Unsited Person", "branch": ""}]
r8 = xlsx_io.parse_paste("2026-11-13\tJohn Lowell\n2026-11-13\tSite B Person\n2026-11-14\tUnsited Person", _mixed)
check("paste: the same date is fine for an site_a and a site_b person", [i for _, i in r8["ok"]] == [1, 9])
check("paste: a person without a branch is reported", len(r8["bad"]) == 1 and "no site" in r8["bad"][0][2])
try:
    xlsx_io.parse_paste("\n".join("2026-01-01\tJohn Lowell" for _ in range(367)), _staff)
    check("paste: more than 366 rows refused", False)
except xlsx_io.TooManyRows:
    check("paste: more than 366 rows refused", True)
check("paste: exactly 366 rows is allowed", len(xlsx_io.parse_paste("\n".join("2026-01-01\tJohn Lowell" for _ in range(366)), _staff)["bad"]) == 365)

blob = xlsx_io.build_xlsx([("2026-11-02", "Monday", "Site A", "John Lowell", "511 22 33 44", ), ("2026-11-03", "Tuesday", "Site B", "=HYPERLINK(1)", "")])
import openpyxl as _ox
_ws = _ox.load_workbook(_io.BytesIO(blob)).active
check("xlsx: header and rows", [c.value for c in _ws[1]] == ["Date", "Day", "Site", "On duty", "Phone"] and _ws.cell(2, 4).value == "John Lowell")
check("xlsx: text starting with = is stored as text, not a formula", _ws.cell(3, 4).data_type == "s")

with app.app_context():
    n_before = main.q1("SELECT COUNT(*) AS n FROM duty")["n"]
imp_text = f"{(date.today() + timedelta(days=200)).strftime('%d.%m.%Y')}\tNew Person\n{(date.today() + timedelta(days=201)).isoformat()}\tSecond Person\n{(date.today() + timedelta(days=202)).isoformat()}\tNobody"
prev = owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": imp_text, "step": "preview"})
check("import: preview shows rows and the problem line, writes nothing", prev.status_code == 200 and "New Person" in prev.get_data(as_text=True) and "Nobody" in prev.get_data(as_text=True))
with app.app_context():
    check("import: preview did not change the schedule", main.q1("SELECT COUNT(*) AS n FROM duty")["n"] == n_before)
owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": imp_text, "step": "confirm"})
with app.app_context():
    check("import: confirm writes only the good rows", main.q1("SELECT COUNT(*) AS n FROM duty")["n"] == n_before + 2)
    check("import: one audit entry", main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'duty.import'")["n"] == 1)
check("import: 367 rows refused (400)", owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": "\n".join("2026-01-01\tx" for _ in range(367)), "step": "preview"}).status_code == 400)
check("import: needs login", anon.post("/admin/schedule/import", data={"text": "x", "step": "preview"}).status_code in (302, 400, 403))

ex = owner.get("/admin/export.xlsx?from=" + (date.today() + timedelta(days=200)).isoformat() + "&to=" + (date.today() + timedelta(days=202)).isoformat())
_ws2 = _ox.load_workbook(_io.BytesIO(ex.data)).active
check("export: period returns a workbook with the assigned days", ex.status_code == 200 and "spreadsheetml" in ex.content_type and _ws2.max_row >= 3 and _ws2.cell(2, 4).value and _ws2.cell(2, 3).value == "Site A")
check("export: week view works", owner.get("/admin/export.xlsx?view=week&d=" + date.today().isoformat()).status_code == 200)
check("export: month view works", owner.get("/admin/export.xlsx?view=month&d=" + date.today().isoformat()).status_code == 200)
check("export: bad dates refused (400)", owner.get("/admin/export.xlsx?from=zzz&to=zzz").status_code == 400)
check("export: reversed or huge period refused (400)", owner.get("/admin/export.xlsx?from=2026-12-01&to=2026-01-01").status_code == 400 and owner.get("/admin/export.xlsx?from=2020-01-01&to=2030-01-01").status_code == 400)
check("export: needs login", anon.get("/admin/export.xlsx").status_code in (302, 401, 403))
check("schedule page: import and export controls present", 'id="dlg-import"' in owner.get("/admin/schedule").get_data(as_text=True) and "export.xlsx" in owner.get("/admin/schedule").get_data(as_text=True))

# --- remove a person from "Staff"
owner.post("/admin/staff/save", data={"_csrf": token(owner, "/admin/staff"), "name": "Person To Delete", "branch": "site_a", "email": "", "phone": ""})
owner.post("/admin/staff/save", data={"_csrf": token(owner, "/admin/staff"), "name": "Person To Keep", "branch": "site_a", "email": "", "phone": ""})
with app.app_context():
    gone = main.q1("SELECT id FROM staff WHERE name = ?", "Person To Delete")["id"]
    stay = main.q1("SELECT id FROM staff WHERE name = ?", "Person To Keep")["id"]
past_d = (date.today() - timedelta(days=5)).isoformat()
fut_d = (date.today() + timedelta(days=300)).isoformat()
fut_d2 = (date.today() + timedelta(days=301)).isoformat()
for dd, ss in ((past_d, gone), (fut_d, gone), (fut_d2, stay)):
    owner.post("/admin/schedule/assign", data={"_csrf": token(owner, "/admin/schedule"), "day": dd, "staff_id": ss})
page = owner.get("/admin/staff").get_data(as_text=True)
check("staff page: delete button with exact day counts in the confirmation", f"/admin/staff/{gone}/delete" in page and "1 day(s)" in page)
check("staff delete: needs login", anon.post(f"/admin/staff/{gone}/delete", data={}).status_code in (302, 400, 403))
check("staff delete: unknown person gives 404", owner.post("/admin/staff/99999/delete", data={"_csrf": token(owner, "/admin/staff")}).status_code == 404)
owner.post(f"/admin/staff/{gone}/delete", data={"_csrf": token(owner, "/admin/staff")})
with app.app_context():
    check("staff delete: the person is gone", main.q1("SELECT 1 FROM staff WHERE id = ?", gone) is None)
    check("staff delete: their duty days are gone (past and future)", main.q1("SELECT COUNT(*) AS n FROM duty WHERE staff_id = ?", gone)["n"] == 0
          and main.q1("SELECT 1 FROM duty WHERE day IN (?, ?)", past_d, fut_d) is None)
    check("staff delete: other people and their days are untouched", main.q1("SELECT 1 FROM staff WHERE id = ?", stay) is not None and main.q1("SELECT staff_id FROM duty WHERE day = ?", fut_d2)["staff_id"] == stay)
    last = main.q1("SELECT detail FROM audit WHERE action = 'staff.delete' ORDER BY id DESC LIMIT 1")
    check("staff delete: audit entry names the person and the counts", last is not None and "Person To Delete" in last["detail"] and "1" in last["detail"])
check("staff delete: the schedule page still renders", owner.get("/admin/schedule").status_code == 200)

# --- Task 5: staff screen (status text, counts, warning, paste import) and users (last login)
_ex = [{"name": "Some One"}]
sp = xlsx_io.parse_staff_paste("Full name\tPhone\nNewcomer\t511 22 33 44\n\nSecond Newcomer\t\r\nThird Newcomer\tabc\n", _ex)
check("staff paste: header skipped, phone optional, bad phone reported with its line", [r[0] for r in sp["ok"]] == ["Newcomer", "Second Newcomer"] and sp["ok"][0][1] == "511 22 33 44" and len(sp["bad"]) == 1 and sp["bad"][0][0] == 5)
sp = xlsx_io.parse_staff_paste("Some One\t\nOther Person\t511111111\nOther Person\t522222222\nNewcomer\t511111111", _ex)
check("staff paste: existing and repeated names are refused", [r[0] for r in sp["ok"]] == ["Other Person", "Newcomer"] and len(sp["bad"]) == 2)
sp = xlsx_io.parse_staff_paste("Old Format\told@example.com\t511000000\tSite B", _ex)
check("staff paste: an old list with an e-mail column still works, the e-mail is dropped", sp["ok"] == [("Old Format", "511000000", "site_b")] and not sp["bad"])
sp = xlsx_io.parse_staff_paste("Site B Newcomer\t\tSite B\nSite A Newcomer\tsite a\n\nBad Site\t\tNorth", _ex)
check("staff paste: optional branch column, a bad value is reported", [(r[0], r[2]) for r in sp["ok"]] == [("Site B Newcomer", "site_b"), ("Site A Newcomer", "site_a")] and len(sp["bad"]) == 1)
sp = xlsx_io.parse_staff_paste("Person One;511000000\nPerson Two;;", _ex)
check("staff paste: semicolon separated rows work", [r[0] for r in sp["ok"]] == ["Person One", "Person Two"])
try:
    xlsx_io.parse_staff_paste("\n".join(f"Person {i}" for i in range(201)), _ex)
    check("staff paste: more than 200 rows refused", False)
except xlsx_io.TooManyRows:
    check("staff paste: more than 200 rows refused", True)
check("staff paste: exactly 200 rows allowed", len(xlsx_io.parse_staff_paste("\n".join(f"Person {i}" for i in range(200)), _ex)["ok"]) == 200)

prev = owner.post("/admin/staff/import", data={"_csrf": token(owner, "/admin/staff"), "step": "preview", "text": "Import Person\t511 00 00 01\nImport Person Two\tbad\n"})
check("staff import: preview lists good and bad rows and writes nothing", prev.status_code == 200 and "Import Person" in prev.get_data(as_text=True) and "E-mail" not in prev.get_data(as_text=True))
with app.app_context():
    check("staff import: preview saved nobody", main.q1("SELECT 1 FROM staff WHERE name = ?", "Import Person") is None)
owner.post("/admin/staff/import", data={"_csrf": token(owner, "/admin/staff"), "step": "confirm", "text": "Import Person\t511 00 00 01\nImport Person Two\tbad\n"})
with app.app_context():
    check("staff import: confirm adds only the good row", main.q1("SELECT phone FROM staff WHERE name = ?", "Import Person")["phone"] == "511 00 00 01" and main.q1("SELECT 1 FROM staff WHERE name = ?", "Import Person Two") is None)
    check("staff import: one audit entry", main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'staff.import'")["n"] == 1)
check("staff import: 201 rows refused (400)", owner.post("/admin/staff/import", data={"_csrf": token(owner, "/admin/staff"), "step": "preview", "text": "\n".join(f"Person {i}" for i in range(201))}).status_code == 400)
check("staff import: needs login", anon.post("/admin/staff/import", data={"text": "x", "step": "preview"}).status_code in (302, 400, 403))

owner.post("/admin/staff/save", data={"_csrf": token(owner, "/admin/staff"), "name": "Inactive Person", "branch": "site_a", "email": "", "phone": ""})
with app.app_context():
    off_id = main.q1("SELECT id FROM staff WHERE name = ?", "Inactive Person")["id"]
fd1, fd2 = [(date.today() + timedelta(days=n)).isoformat() for n in (400, 401)]
for dd in (fd1, fd2):
    owner.post("/admin/schedule/assign", data={"_csrf": token(owner, "/admin/schedule"), "day": dd, "staff_id": off_id})
page = owner.get("/admin/staff").get_data(as_text=True)
check("staff page: plain-text status, no on/off switch", 'class="switch' not in page and "Active" in page and "Deactivate" in page)
check("staff page: no warning while the person is active", "Inactive but" not in page)
owner.post(f"/admin/staff/{off_id}/toggle", data={"_csrf": token(owner, "/admin/staff")})
page = owner.get("/admin/staff").get_data(as_text=True)
check("staff page: deactivated person with future days gets an exact warning", "Inactive but has 2 upcoming day(s)" in page and "Activate" in page)
check("staff page: the future days stay visible in the schedule", duty_of(fd1) == off_id)
users_html = owner.get("/admin/users").get_data(as_text=True)
with app.app_context():
    ll = main.q1("SELECT last_login FROM users WHERE username = 'owner'")["last_login"]
check("users page: last login is recorded and shown", ll is not None and "Last login" in users_html)

# --- login visibility for the owner: last IP, host names (written by the VM's reverse-DNS helper), history
with app.app_context():
    _row = main.q1("SELECT last_ip FROM users WHERE username = 'owner'")
check("login: the client IP of the last login is stored", _row["last_ip"] == "127.0.0.1")
_page = owner.get("/admin/users").get_data(as_text=True)
check("users page: IP column and the login history are shown", "127.0.0.1" in _page and 'id="login-history"' in _page)
check("users page: failed attempts appear in the history", "Failed attempt" in _page)
_hp = os.path.join(os.environ["DATA_DIR"], "hostnames.json")
def _write_hosts(text, bump):
    with open(_hp, "w", encoding="utf-8") as f:
        f.write(text)
    os.utime(_hp, (time.time() + bump, time.time() + bump))      # a new modification time, whatever the clock resolution
_write_hosts('{"127.0.0.1": "it-pc.example.local"}', 5)
check("users page: host name from hostnames.json is shown", "it-pc.example.local" in owner.get("/admin/users").get_data(as_text=True))
_write_hosts('{"127.0.0.1": "<b>x</b>"}', 10)
_page = owner.get("/admin/users").get_data(as_text=True)
check("users page: host names are HTML-escaped", "&lt;b&gt;x&lt;/b&gt;" in _page and "<b>x</b>" not in _page)
_write_hosts("this is not json", 15)
check("users page: a broken hostnames.json is ignored", owner.get("/admin/users").status_code == 200)
os.remove(_hp)
check("users page: works without hostnames.json", owner.get("/admin/users").status_code == 200)
with app.app_context():
    _now = main.iso(main.now_utc())
    _old = main.iso(main.now_utc() - timedelta(days=90))
    for _ts, _act, _ip in ((_now, "login", "10.1.1.1"), (_now, "login.failed", "10.1.1.2"), (_now, "duty.assign", "10.9.9.9"), (_old, "login", "10.2.2.2")):
        main.run("INSERT INTO audit(ts, actor, action, detail, ip) VALUES (?, 'x', ?, '', ?)", _ts, _act, _ip)
_res = runner.invoke(args=["login-ips"])
_ips = _res.output.split()
check("CLI login-ips: addresses of login events, each once", _res.exit_code == 0 and "10.1.1.1" in _ips and "10.1.1.2" in _ips and len(_ips) == len(set(_ips)))
check("CLI login-ips: no other events, nothing older than 60 days", "10.9.9.9" not in _ips and "10.2.2.2" not in _ips)

# --- owner creates an admin: temp password shown once, must change at first login
r = owner.post("/admin/users/create", data={"_csrf": token(owner, "/admin/users"), "username": "nino", "display_name": "Nina"})
temp = re.search(r'id="secret">([^<]+)<', r.get_data(as_text=True))
check("temporary password shown", r.status_code == 200 and temp is not None)
temp = temp.group(1)

admin = app.test_client()
check("new admin can log in with temp password", login(admin, "nino", temp).status_code == 302)
check("forced to change password first", admin.get("/admin/schedule").status_code == 302 and "settings" in admin.get("/admin/schedule").headers["Location"])
NEW_PW = "Another-Pass-9"
r = admin.post("/admin/settings", data={"_csrf": token(admin, "/admin/settings"), "current": temp, "new": "short", "again": "short"})
check("short password refused (min 7)", "at least 7" in r.get_data(as_text=True))
r = admin.post("/admin/settings", data={"_csrf": token(admin, "/admin/settings"), "current": temp, "new": "alllowercasepassword", "again": "alllowercasepassword"})
check("long but simple password refused (complexity)", "upper-case letter" in r.get_data(as_text=True))
r = admin.post("/admin/settings", data={"_csrf": token(admin, "/admin/settings"), "current": temp, "new": "Ab1!xyz", "again": "Ab1!xyz"})
check("7-character complex password is accepted", "Password changed" in r.get_data(as_text=True))
temp = "Ab1!xyz"
r = admin.post("/admin/settings", data={"_csrf": token(admin, "/admin/settings"), "current": temp, "new": NEW_PW, "again": NEW_PW})
check("password changed", "Password changed" in r.get_data(as_text=True))
check("admin reaches schedule now", admin.get("/admin/schedule").status_code == 200)
check("admin cannot open users page (403)", admin.get("/admin/users").status_code == 403)
with app.app_context():
    owner_row = main.q1("SELECT id FROM users WHERE username = 'owner'")["id"]
check("admin cannot reset another user's password (403)", admin.post(f"/admin/users/{owner_row}/reset", data={"_csrf": token(admin, "/admin/settings")}).status_code == 403)
check("admin cannot create users (403)", admin.post("/admin/users/create", data={"_csrf": token(admin, "/admin/settings"), "username": "evil", "display_name": "x"}).status_code == 403)
check("admin cannot unlock a locked computer (403)", admin.post("/admin/locks/unlock", data={"_csrf": token(admin, "/admin/settings"), "username": "nino", "ip": "10.1.2.3"}).status_code == 403)
check("admin cannot unlock or disable users (403)", admin.post(f"/admin/users/{owner_row}/toggle", data={"_csrf": token(admin, "/admin/settings")}).status_code == 403)

# --- Task 6: site settings and backups (owner only)
adm_settings = admin.get("/admin/settings").get_data(as_text=True)
check("settings: non-owner sees no site or backup controls", 'id="site-form"' not in adm_settings and 'id="backup-box"' not in adm_settings)
check("settings: non-owner cannot save site texts (403)", admin.post("/admin/settings/site", data={"_csrf": token(admin, "/admin/settings"), "title": "x"}).status_code == 403)
check("backup: non-owner cannot create or download (403)", admin.post("/admin/backup", data={"_csrf": token(admin, "/admin/settings")}).status_code == 403 and admin.get("/admin/backup/download").status_code == 403)
own_settings = owner.get("/admin/settings").get_data(as_text=True)
check("settings: owner sees site form, backup box and the read-only info block", 'id="site-form"' in own_settings and 'id="backup-box"' in own_settings and 'id="info-box"' in own_settings and "8" in own_settings)
site = {"title": "Test title", "footer": "Test footer", "contact_phone": "511 22 33 44"}
r = owner.post("/admin/settings/site", data={"_csrf": token(owner, "/admin/settings"), **site})
check("settings: site texts saved and shown again", r.status_code in (200, 302) and "Test title" in owner.get("/admin/settings").get_data(as_text=True))
pub = anon.get("/").get_data(as_text=True)
check("settings: new title and footer appear on the public page", "Test title" in pub and "Test footer" in pub)
check("settings: empty title refused (400)", owner.post("/admin/settings/site", data={"_csrf": token(owner, "/admin/settings"), **{**site, "title": "  "}}).status_code == 400)
check("settings: too-long title refused (400)", owner.post("/admin/settings/site", data={"_csrf": token(owner, "/admin/settings"), **{**site, "title": "a" * 81}}).status_code == 400)
with app.app_context():
    check("settings: change written to the audit log", main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'settings.site'")["n"] >= 1)
bk_dir = os.path.join(os.environ["DATA_DIR"], "backups")
n0 = len(os.listdir(bk_dir)) if os.path.isdir(bk_dir) else 0
owner.post("/admin/backup", data={"_csrf": token(owner, "/admin/settings")})
check("backup: 'create now' writes a new file", len(os.listdir(bk_dir)) >= n0 + 1)
with app.app_context():
    check("backup: audit entry", main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'backup.create'")["n"] >= 1)
check("backup: settings page shows the last backup", "Last backup" in owner.get("/admin/settings").get_data(as_text=True))
dl = owner.get("/admin/backup/download")
check("backup: download is an attachment holding a real SQLite file", dl.status_code == 200 and "attachment" in dl.headers.get("Content-Disposition", "") and dl.data[:15] == b"SQLite format 3")
check("backup: download needs login", anon.get("/admin/backup/download").status_code in (302, 401, 403))

# --- owner resets the admin's password (forgotten password flow)
with app.app_context():
    nino_id = main.q1("SELECT id FROM users WHERE username = 'nino'")["id"]
r = owner.post(f"/admin/users/{nino_id}/reset", data={"_csrf": token(owner, "/admin/users")})
temp2 = re.search(r'id="secret">([^<]+)<', r.get_data(as_text=True)).group(1)
check("reset ends the admin's existing session", admin.get("/admin/schedule").status_code == 302 and "login" in admin.get("/admin/schedule").headers["Location"])
check("old password no longer works", login(app.test_client(), "nino", NEW_PW).status_code == 401)
check("reset temp password works", login(app.test_client(), "nino", temp2).status_code == 302)
with app.app_context():
    owner_id = main.q1("SELECT id FROM users WHERE username = 'owner'")["id"]
check("owner cannot reset own password via panel", owner.post(f"/admin/users/{owner_id}/reset", data={"_csrf": token(owner, "/admin/users")}).status_code == 400)

# --- lockout after 5 failures, then owner unlocks
victim = app.test_client()
for _ in range(5):
    login(victim, "nino", "definitely-wrong-pass")
check("locked: correct password refused", login(victim, "nino", temp2).status_code == 401)
owner.post(f"/admin/users/{nino_id}/unlock", data={"_csrf": token(owner, "/admin/users")})
check("unlocked by owner: login works", login(app.test_client(), "nino", temp2).status_code == 302)

# --- lockout is per user AND client address; the owner can unlock one computer or the whole account
def client_from(ip):
    c = app.test_client()
    c.environ_base["REMOTE_ADDR"] = ip
    return c
def lock_rows(where="1=1", *args):
    with app.app_context():
        return main.q1(f"SELECT COUNT(*) AS n FROM login_locks WHERE {where}", *args)["n"]
for _ in range(5):
    login(client_from("10.9.9.1"), "nino", "wrong-wrong-pass")
_r = login(client_from("10.9.9.1"), "nino", temp2)
check("lock: from the locked address even the right password is refused, with a clear message", _r.status_code == 401 and "locked" in _r.get_data(as_text=True))
with app.app_context():
    _lk = main.q1("SELECT ip FROM audit WHERE action = 'login.locked' AND actor = 'nino' ORDER BY id DESC LIMIT 1")
check("lock: the audit entry records the address that was locked", _lk is not None and _lk["ip"] == "10.9.9.1")
check("lock: the same user from ANOTHER address is not affected", login(client_from("10.9.9.2"), "nino", temp2).status_code == 302)
check("lock: another user from the locked address is not affected", login(client_from("10.9.9.1"), "owner", OWNER_PW).status_code == 302)
_before = lock_rows()
for _ in range(6):
    _g = login(client_from("10.9.9.50"), "no-such-user", "whatever-pass")
check("lock: unknown user names are refused but leave no lock rows behind", _g.status_code == 401 and lock_rows() == _before)
_page = owner.get("/admin/users").get_data(as_text=True)
check("users page: locked computers are listed with an unlock button", 'id="locked-list"' in _page and "10.9.9.1" in _page and "/admin/locks/unlock" in _page)
for ip in ("10.9.9.3", "10.9.9.4"):
    for _ in range(5):
        login(client_from(ip), "nino", "wrong-wrong-pass")
owner.post("/admin/locks/unlock", data={"_csrf": token(owner, "/admin/users"), "username": "nino", "ip": "10.9.9.3"})
check("unlock one computer: that computer can log in again", login(client_from("10.9.9.3"), "nino", temp2).status_code == 302)
check("unlock one computer: another locked computer stays locked", login(client_from("10.9.9.4"), "nino", temp2).status_code == 401)
with app.app_context():
    _ua = main.q1("SELECT detail FROM audit WHERE action = 'user.unlock' ORDER BY id DESC LIMIT 1")
check("unlock: written to the audit log with the address", _ua is not None and "10.9.9.3" in _ua["detail"])
owner.post(f"/admin/users/{nino_id}/unlock", data={"_csrf": token(owner, "/admin/users")})
check("unlock the account: every locked computer works again", login(client_from("10.9.9.4"), "nino", temp2).status_code == 302 and login(client_from("10.9.9.1"), "nino", temp2).status_code == 302)
for _ in range(4):
    login(client_from("10.9.9.6"), "nino", "wrong-wrong-pass")
login(client_from("10.9.9.6"), "nino", temp2)
for _ in range(4):
    login(client_from("10.9.9.6"), "nino", "wrong-wrong-pass")
check("lock: a successful login resets the failure counter", login(client_from("10.9.9.6"), "nino", temp2).status_code == 302)
with app.app_context():
    main.run("INSERT INTO login_locks(username, ip, failed, last_fail, locked_until) VALUES ('nino', '10.9.9.7', 4, ?, NULL)", main.iso(main.now_utc() - timedelta(minutes=20)))
login(client_from("10.9.9.7"), "nino", "wrong-wrong-pass")
check("lock: failures older than 15 minutes do not add up", login(client_from("10.9.9.7"), "nino", temp2).status_code == 302)
with app.app_context():
    main.run("INSERT INTO login_locks(username, ip, failed, last_fail, locked_until) VALUES ('nino', '10.9.9.9', 0, ?, ?)", main.iso(main.now_utc() - timedelta(minutes=30)), main.iso(main.now_utc() - timedelta(minutes=1)))
check("lock: an expired lock no longer blocks", login(client_from("10.9.9.9"), "nino", temp2).status_code == 302)
for _ in range(5):
    login(client_from("10.9.9.8"), "nino", "wrong-wrong-pass")
_cli = runner.invoke(args=["unlock-user"], input="nino\n")
check("CLI unlock-user clears the locks of an account", _cli.exit_code == 0 and login(client_from("10.9.9.8"), "nino", temp2).status_code == 302)

# --- break-glass CLI reset for the owner
res = runner.invoke(args=["reset-password"], input="owner\n")
cli_temp = re.search(r"\n\s+(\S{16})\s*$", res.output.strip("\n") + "\n")
check("CLI reset-password prints a temp password", res.exit_code == 0 and cli_temp is not None)
check("CLI reset ended the owner's old session", owner.get("/admin/schedule").status_code == 302 and "login" in owner.get("/admin/schedule").headers["Location"])
owner = app.test_client()
check("owner can log in with CLI temp password", login(owner, "owner", cli_temp.group(1)).status_code == 302)
check("...and is forced to choose a new password", "settings" in owner.get("/admin/schedule").headers["Location"])
OWNER_PW2 = "Brand-New-Owner-3"
owner.post("/admin/settings", data={"_csrf": token(owner, "/admin/settings"), "current": cli_temp.group(1), "new": OWNER_PW2, "again": OWNER_PW2})
check("owner is back in the schedule", owner.get("/admin/schedule").status_code == 200)

# --- audit + backup
log = owner.get("/admin/audit").get_data(as_text=True)
check("audit has password reset entries", "user.password_reset" in log and "duty.assign" in log)
res = runner.invoke(args=["backup"])
check("CLI backup writes a file", res.exit_code == 0 and os.path.exists(res.output.strip()))

# --- Task 6: audit filters, IP, export
with app.app_context():
    ips = main.q1("SELECT ip FROM audit WHERE action = 'duty.swap' LIMIT 1")["ip"]
check("audit: the client IP is stored with new entries", ips == "127.0.0.1")
page = owner.get("/admin/audit").get_data(as_text=True)
check("audit: IP column shown", "127.0.0.1" in page)
page = owner.get("/admin/audit?user=nino").get_data(as_text=True)
check("audit: filter by user", "<code>password.change</code>" in page and "<code>duty.swap</code>" not in page)
page = owner.get("/admin/audit?action=duty.swap").get_data(as_text=True)
check("audit: filter by action", "<code>duty.swap</code>" in page and "<code>staff.add</code>" not in page)
page = owner.get("/admin/audit?q=" + "Person To Delete").get_data(as_text=True)
check("audit: search in the details", "<code>staff.add</code>" in page and "<code>duty.swap</code>" not in page)
_t = date.today().isoformat()
check("audit: period that contains today finds entries", "<code>duty.swap</code>" in owner.get(f"/admin/audit?from={_t}&to={_t}").get_data(as_text=True))
check("audit: period in the past finds nothing and says so", "No entries found" in owner.get("/admin/audit?from=2000-01-01&to=2000-01-02").get_data(as_text=True))
check("audit: filters combine", "<code>duty.swap</code>" not in owner.get("/admin/audit?user=nino&action=duty.swap").get_data(as_text=True))
check("audit: malformed date refused (400)", owner.get("/admin/audit?from=zzz").status_code == 400)
ax = owner.get("/admin/audit.xlsx?action=duty.swap")
_wa = _ox.load_workbook(_io.BytesIO(ax.data)).active
with app.app_context():
    n_swap = main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'duty.swap'")["n"]
check("audit export: header with IP column", ax.status_code == 200 and [c.value for c in _wa[1]] == ["Time", "User", "Action", "Details", "IP"])
check("audit export: only the filtered rows", _wa.max_row == n_swap + 1 and all(_wa.cell(i, 3).value == "duty.swap" for i in range(2, _wa.max_row + 1)))
check("audit export: needs login", anon.get("/admin/audit.xlsx").status_code in (302, 401, 403))
check("audit page: shows how many entries match", f"{n_swap}" in owner.get("/admin/audit?action=duty.swap").get_data(as_text=True))

# ================= Site A / Site B model =================
def post_owner(path, **data):
    return owner.post(path, data={"_csrf": token(owner, "/admin/staff"), **data})
def new_person(name, branch):
    post_owner("/admin/staff/save", name=name, branch=branch, email="", phone="")
    with app.app_context():
        return main.q1("SELECT id FROM staff WHERE name = ?", name)["id"]
def slot(dd, branch):
    with app.app_context():
        r = main.q1("SELECT staff_id FROM duty WHERE day = ? AND branch = ?", dd, branch)
        return r["staff_id"] if r else None
def has(dd, branch, pid):
    with app.app_context():
        return main.q1("SELECT 1 FROM duty WHERE day = ? AND branch = ? AND staff_id = ?", dd, branch, pid) is not None
def branch_of(pid):
    with app.app_context():
        return main.q1("SELECT branch FROM staff WHERE id = ?", pid)["branch"]

e1 = new_person("Site A One", "site_a")
w1 = new_person("Site B One", "site_b")
n1 = new_person("Unsited One", "")
dd1 = (date.today() + timedelta(days=500)).isoformat()
post_owner("/admin/schedule/assign", day=dd1, staff_id=e1)
post_owner("/admin/schedule/assign", day=dd1, staff_id=w1)
check("two slots: one day holds an site_a AND a site_b person", slot(dd1, "site_a") == e1 and slot(dd1, "site_b") == w1)
check("assign: a person goes to THEIR branch's slot, whatever the form says", post_owner("/admin/schedule/assign", day=dd1, staff_id=e1, branch="site_b").status_code == 400)
check("assign: a person without a branch cannot be placed (400)", post_owner("/admin/schedule/assign", day=dd1, staff_id=n1).status_code == 400)
check("assign: a missing day is refused (400)", post_owner("/admin/schedule/assign", day="", staff_id=e1).status_code == 400)
post_owner("/admin/schedule/assign", day=dd1, staff_id="", branch="site_a")
check("clear: empties only the chosen branch's slot", slot(dd1, "site_a") is None and slot(dd1, "site_b") == w1)
check("clear: an unknown branch is refused (400)", post_owner("/admin/schedule/assign", day=dd1, staff_id="", branch="north").status_code == 400)
post_owner("/admin/schedule/bulk", days=dd1, staff_id="", branch="site_b")
check("bulk clear: per branch", slot(dd1, "site_b") is None)
check("bulk clear: needs a branch (400)", post_owner("/admin/schedule/bulk", days=dd1, staff_id="").status_code == 400)
w2 = new_person("Site B Two", "site_b")
dd2, dd3 = (date.today() + timedelta(days=501)).isoformat(), (date.today() + timedelta(days=502)).isoformat()
post_owner("/admin/schedule/assign", day=dd2, staff_id=w1)
post_owner("/admin/schedule/assign", day=dd3, staff_id=w2)
post_owner("/admin/schedule/swap", a=dd2, b=dd3, branch="site_b")
check("swap: works inside the site_b line", slot(dd2, "site_b") == w2 and slot(dd3, "site_b") == w1)
post_owner("/admin/schedule/assign", day=dd2, staff_id=e1)
post_owner("/admin/schedule/swap", a=dd2, b=dd3, branch="site_b")
check("swap: leaves the other branch's line alone", slot(dd2, "site_a") == e1)

# the public page and the API show both branches
dt = date.today().isoformat()
post_owner("/admin/schedule/assign", day=dt, staff_id="", branch="site_a")
post_owner("/admin/schedule/assign", day=dt, staff_id="", branch="site_b")
post_owner("/admin/schedule/assign", day=dt, staff_id=e1)
post_owner("/admin/schedule/assign", day=dt, staff_id=w1)
pub = anon.get("/").get_data(as_text=True)
check("public: two cards with the branch labels", "Site A" in pub and "Site B" in pub and pub.count('class="duty-card"') == 2)
check("public: both on-duty names are on the cards", "Site A One" in pub and "Site B One" in pub)
check("public: week lists have the Site A and Site B column heads", pub.count('class="wk-head"') == 2 and "Site A</span><span>Site B" in pub)
check("public: today/tomorrow tags sit in the date cell", re.search(r'wk-daycell.*?<span class="pill soft">Today</span>.*?</span>\s*<span class="wk-who site_a"', pub, re.S) is not None)
api = anon.get("/api/today").get_json()["on_duty"]
check("api/today: site_a and site_b", [p["name"] for p in api["site_a"]] == ["Site A One"] and [p["name"] for p in api["site_b"]] == ["Site B One"])
post_owner("/admin/schedule/assign", day=dt, staff_id="", branch="site_b")
pub = anon.get("/").get_data(as_text=True)
check("public: an empty slot shows the amber 'not assigned' state", pub.count("Not assigned") >= 1 and 'class="empty-state"' in pub)
check("api/today: an empty slot is an empty list", anon.get("/api/today").get_json()["on_duty"]["site_b"] == [])

# admin schedule page
page = owner.get("/admin/schedule").get_data(as_text=True)
check("schedule page: two lines per day, palette grouped by branch", 'class="ln e' in page and 'class="ln w' in page and 'gh e">Site A' in page and "No site chosen" in page)
check("schedule page: per-branch clear buttons and no weekend checkbox", 'id="bulk-clear-site_a"' in page and 'id="bulk-clear-site_b"' in page and 'name="weekends"' not in page)

# the branch switch on the staff page
page = owner.get("/admin/staff").get_data(as_text=True)
check("staff page: toggle switch with both labels and the confirmation dialog", 'class="swx"' in page and 'id="dlg-branch"' in page and "Change of site" in page and "No site chosen" in page)
sw = new_person("Switch Person", "site_a")
past_d = (date.today() - timedelta(days=800)).isoformat()
fut1, fut2 = (date.today() + timedelta(days=600)).isoformat(), (date.today() + timedelta(days=601)).isoformat()
with app.app_context():
    main.run("INSERT INTO duty(day, branch, staff_id) VALUES (?, 'site_a', ?)", past_d, sw)
for dd in (dt, fut1, fut2):
    post_owner("/admin/schedule/assign", day=dd, staff_id=sw)
before = audit_count("staff.branch")
check("branch switch: invalid branch refused (400)", post_owner(f"/admin/staff/{sw}/branch", branch="north").status_code == 400)
check("branch switch: unknown person gives 404", post_owner("/admin/staff/99999/branch", branch="site_b").status_code == 404)
post_owner(f"/admin/staff/{sw}/branch", branch="site_a")
check("branch switch: same branch changes nothing and writes no audit entry", branch_of(sw) == "site_a" and audit_count("staff.branch") == before and slot(fut1, "site_a") == sw)
post_owner(f"/admin/staff/{sw}/branch", branch="site_b")
check("branch switch: the person is now site_b", branch_of(sw) == "site_b")
check("branch switch: their old-branch days from today on are empty, not moved", not has(dt, "site_a", sw) and not has(fut1, "site_a", sw) and not has(fut2, "site_a", sw) and not has(fut1, "site_b", sw) and not has(dt, "site_b", sw))
check("branch switch: days before today stay as history", slot(past_d, "site_a") == sw)
with app.app_context():
    last = main.q1("SELECT detail FROM audit WHERE action = 'staff.branch' ORDER BY id DESC LIMIT 1")["detail"]
check("branch switch: audited with old -> new and the count", audit_count("staff.branch") == before + 1 and "Site A" in last and "Site B" in last and "3" in last)
adm2 = app.test_client()
login(adm2, "nino", temp2)
adm2.post("/admin/settings", data={"_csrf": token(adm2, "/admin/settings"), "current": temp2, "new": "Nino-Pass-77", "again": "Nino-Pass-77"})
check("branch switch: any administrator may switch a person", adm2.post(f"/admin/staff/{sw}/branch", data={"_csrf": token(adm2, "/admin/staff"), "branch": "site_a"}).status_code == 302 and branch_of(sw) == "site_a")
check("branch switch: needs login", anon.post(f"/admin/staff/{sw}/branch", data={"branch": "site_b"}).status_code in (302, 400, 403))
# a person without a branch adopts the days that were assigned before branches existed
with app.app_context():
    main.run("INSERT INTO duty(day, branch, staff_id) VALUES (?, '', ?)", fut1, n1)
post_owner(f"/admin/staff/{n1}/branch", branch="site_b")
check("legacy days: the person's old 'no branch' days move into their new branch", slot(fut1, "") is None and slot(fut1, "site_b") == n1)
with app.app_context():
    main.run("INSERT INTO duty(day, branch, staff_id) VALUES (?, '', ?)", fut2, e1)
check("schedule page: legacy line renders", 'class="ln n"' in owner.get("/admin/schedule?d=" + fut2).get_data(as_text=True))
check("legacy days: clearing the 'no branch' slot works", post_owner("/admin/schedule/assign", day=fut2, staff_id="", branch="").status_code == 302 and slot(fut2, "") is None)

# Excel import uses the person's branch; export has the branch column
txt = f"{(date.today() + timedelta(days=700)).isoformat()}\tSite B One\n{(date.today() + timedelta(days=701)).isoformat()}\tUnsited One"
prev = owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": txt, "step": "preview"}).get_data(as_text=True)
check("import preview: shows the person's branch", "Site B" in prev)
n1b = new_person("Unsited Two", "")
txt2 = f"{(date.today() + timedelta(days=702)).isoformat()}\tUnsited Two"
prev = owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": txt2, "step": "preview"}).get_data(as_text=True)
check("import preview: a person with no branch is reported", "has no site" in prev)
owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "text": txt, "step": "confirm"})
check("import confirm: the person lands in their own branch's slot", slot((date.today() + timedelta(days=700)).isoformat(), "site_b") == w1)

# copy last week copies both slots
mon = date.today() - timedelta(days=date.today().weekday())
src_day = (mon - timedelta(days=7)).isoformat()
dst_day = mon.isoformat()
post_owner("/admin/schedule/assign", day=src_day, staff_id=e1)
post_owner("/admin/schedule/assign", day=src_day, staff_id=w1)
post_owner("/admin/schedule/assign", day=dst_day, staff_id="", branch="site_a")
post_owner("/admin/schedule/assign", day=dst_day, staff_id="", branch="site_b")
post_owner("/admin/schedule/copy-week", d=dst_day)
check("copy week: both slots are copied", slot(dst_day, "site_a") == e1 and slot(dst_day, "site_b") == w1)

# holidays: colouring and the owner's overrides
hd = date.today() + timedelta(days=2)
with app.app_context():
    base = main.holiday_map(date(2027, 1, 1), date(2027, 1, 31))
check("holidays: map has the automatic days for a month", date(2027, 1, 7) in base and date(2027, 1, 19) in base)
check("holidays: add is owner only", adm2.post("/admin/holidays/add", data={"_csrf": token(adm2, "/admin/settings"), "day": hd.isoformat(), "name": "x"}).status_code == 403)
post_owner("/admin/holidays/add", day=hd.isoformat(), name="Test holiday")
pub = anon.get("/").get_data(as_text=True)
check("holidays: a day added by the owner is marked on the public page like a weekend", "Test holiday" in pub and 'wk-row we ho' in pub)
check("holidays: ...and in the admin calendar", "Test holiday" in owner.get("/admin/schedule?d=" + hd.isoformat()).get_data(as_text=True))
check("holidays: a bad date is refused (400)", post_owner("/admin/holidays/add", day="zzz", name="x").status_code == 400)
post_owner("/admin/holidays/delete", day=hd.isoformat())
check("holidays: a custom day can be removed again", "Test holiday" not in anon.get("/").get_data(as_text=True))
post_owner("/admin/holidays/toggle", day="2027-01-07", off="1")
with app.app_context():
    off_map = main.holiday_map(date(2027, 1, 1), date(2027, 1, 31))
check("holidays: an automatic day can be switched off", date(2027, 1, 7) not in off_map and date(2027, 1, 19) in off_map)
post_owner("/admin/holidays/toggle", day="2027-01-07", off="0")
with app.app_context():
    check("holidays: ...and back on", date(2027, 1, 7) in main.holiday_map(date(2027, 1, 1), date(2027, 1, 31)))
check("holidays: only automatic days can be toggled (400)", post_owner("/admin/holidays/toggle", day="2027-02-10", off="1").status_code == 400)
check("holidays: settings page has the holidays block for the owner only", 'id="holidays"' in owner.get("/admin/settings").get_data(as_text=True) and 'id="holidays"' not in adm2.get("/admin/settings").get_data(as_text=True))

# seed-demo created both branches
with app.app_context():
    check("seed-demo: people in both branches", main.q1("SELECT COUNT(DISTINCT branch) AS n FROM staff WHERE branch IN ('site_a','site_b')")["n"] == 2)

# audit deletion (owner only, leaves no trace)
def audit_total():
    with app.app_context():
        return main.q1("SELECT COUNT(*) AS n FROM audit")["n"]
check("audit delete: the panel is there for the owner only", "audit/delete" in owner.get("/admin/audit").get_data(as_text=True) and "audit/delete" not in adm2.get("/admin/audit").get_data(as_text=True))
check("audit delete: administrators are refused (403)", adm2.post("/admin/audit/delete", data={"_csrf": token(adm2, "/admin/settings"), "mode": "all", "word": "delete"}).status_code == 403)
check("audit delete: anonymous visitors are refused", anon.post("/admin/audit/delete", data={"mode": "all", "word": "delete"}).status_code in (302, 400, 403))
with app.app_context():
    ids = [r["id"] for r in main.q("SELECT id FROM audit WHERE action = 'duty.swap' ORDER BY id LIMIT 2")]
t0 = audit_total()
owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "selected", "ids": [str(i) for i in ids]})
check("audit delete: selected rows are gone and NOTHING new was written about it", audit_total() == t0 - len(ids))
t0 = audit_total()
check("audit delete: 'filtered' without any filter is refused (400)", owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "filtered"}).status_code == 400 and audit_total() == t0)
with app.app_context():
    n_sw = main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'duty.assign'")["n"]
owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "filtered", "action": "duty.assign"})
with app.app_context():
    check("audit delete: the filter's rows are gone, other rows stay, no trace", main.q1("SELECT COUNT(*) AS n FROM audit WHERE action = 'duty.assign'")["n"] == 0 and audit_total() == t0 - n_sw)
check("audit delete: 'all' needs the typed word (400 otherwise, nothing deleted)", owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "all", "word": "no"}).status_code == 400 and audit_total() == t0 - n_sw)
check("audit delete: unknown mode refused (400)", owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "x"}).status_code == 400)
owner.post("/admin/audit/delete", data={"_csrf": token(owner, "/admin/audit"), "mode": "all", "word": "delete"})
check("audit delete: everything gone and the log stays empty (no trace of the deletion)", audit_total() == 0 and "audit/delete" in owner.get("/admin/audit").get_data(as_text=True))
check("audit delete: the owner is still logged in afterwards", owner.get("/admin/schedule").status_code == 200)


# ================= several people on duty =================
e2 = new_person("Site A Two", "site_a")
w3 = new_person("Site B Three", "site_b")
dm = (date.today() + timedelta(days=950)).isoformat()
post_owner("/admin/schedule/assign", day=dm, staff_id=e1)
post_owner("/admin/schedule/assign", day=dm, staff_id=e2)
check("multi: assigning a second person ADDS them, nobody is replaced", has(dm, "site_a", e1) and has(dm, "site_a", e2))
post_owner("/admin/schedule/assign", day=dm, staff_id=e2)
with app.app_context():
    check("multi: assigning the same person again changes nothing", main.q1("SELECT COUNT(*) AS n FROM duty WHERE day = ? AND branch = 'site_a'", dm)["n"] == 2)
post_owner("/admin/schedule/assign", day=dm, staff_id=e2, branch="site_a", remove="1")
check("multi: remove takes one person off and leaves the other", has(dm, "site_a", e1) and not has(dm, "site_a", e2))
post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_a", staff_ids=[str(e1), str(e2)])
check("multi: the tick list sets exactly the ticked people", has(dm, "site_a", e1) and has(dm, "site_a", e2))
post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_a", staff_ids=[str(e2)])
check("multi: unticking removes a person", not has(dm, "site_a", e1) and has(dm, "site_a", e2))
post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_a", staff_ids=[])
check("multi: ticking nobody empties the slot", slot(dm, "site_a") is None)
check("multi: a person of the other branch in the tick list is refused (400)", post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_a", staff_ids=[str(w1)]).status_code == 400)
check("multi: a bad branch in set mode is refused (400)", post_owner("/admin/schedule/assign", day=dm, mode="set", branch="north", staff_ids=[]).status_code == 400)
post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_a", staff_ids=[str(e1), str(e2)])
post_owner("/admin/schedule/assign", day=dm, mode="set", branch="site_b", staff_ids=[str(w1), str(w3)])
d2m = (date.today() + timedelta(days=951)).isoformat()
post_owner("/admin/schedule/swap", a=dm, b=d2m, branch="site_a")
check("multi: swap moves the whole group of one branch", has(d2m, "site_a", e1) and has(d2m, "site_a", e2) and not has(dm, "site_a", e1) and has(dm, "site_b", w1) and has(dm, "site_b", w3))
# the schedule page: one button per person, and a tick list in the dialog
page = owner.get("/admin/schedule?d=" + dm).get_data(as_text=True)
check("multi: the calendar shows one line per person", page.count('data-branch="site_b" data-staff="%d"' % w1) >= 1 and page.count('data-branch="site_b" data-staff="%d"' % w3) >= 1)
check("multi: the dialog has tick boxes for the branch's people", 'name="staff_ids"' in page and 'id="assign-list"' in page and 'name="mode" value="set"' in page)
# public page: every person in a green box with the phone below, no e-mail anywhere
with app.app_context():
    main.run("UPDATE staff SET phone = '555 0123', email = 'secret.mail@example.com' WHERE id IN (?, ?)", e1, e2)
post_owner("/admin/schedule/assign", day=dt, staff_id="", branch="site_a")
post_owner("/admin/schedule/assign", day=dt, mode="set", branch="site_a", staff_ids=[str(e1), str(e2)])
pub = anon.get("/").get_data(as_text=True)
card = pub.split("Site A</span>")[1].split("</section>")[0]
check("public: two people on duty give two green boxes in the same card", card.count('class="contact-box person"') == 2 and "Site A One" in card and "Site A Two" in card)
check("public: each box has the name and the phone below it", card.count("555 0123") == 2 and card.index("555 0123") > card.index("Site A One"))
check("public: no e-mail anywhere on the public page or in the API", "@" not in pub.replace("@media", "") and "secret.mail" not in anon.get("/api/today").get_data(as_text=True))
check("public: the week list shows both names in the Site A column", pub.count('class="nm"') >= 2)
check("admin: the staff page shows no e-mail", "secret.mail" not in owner.get("/admin/staff").get_data(as_text=True) and 'name="email"' not in owner.get("/admin/staff").get_data(as_text=True))
check("admin: settings has no e-mail field", 'name="contact_email"' not in owner.get("/admin/settings").get_data(as_text=True))
exm = owner.get("/admin/export.xlsx?from=" + dm + "&to=" + d2m)
_wm = _ox.load_workbook(_io.BytesIO(exm.data)).active
check("multi: the export has one row per person and no e-mail column", _wm.max_row >= 5 and [c.value for c in _wm[1]][-1] == "Phone")
# the import adds several people to one day
dimp = (date.today() + timedelta(days=960)).isoformat()
owner.post("/admin/schedule/import", data={"_csrf": token(owner, "/admin/schedule"), "step": "confirm", "text": f"{dimp}\tSite A One\n{dimp}\tSite A Two"})
check("multi: the list import puts two people on one day", has(dimp, "site_a", e1) and has(dimp, "site_a", e2))

# ================= branch administrators and own-only log =================
r = owner.post("/admin/users/create", data={"_csrf": token(owner, "/admin/users"), "username": "site_a.admin", "display_name": "Site A admin", "branch": "site_a"})
tmp_e = re.search(r'id="secret">([^<]+)<', r.get_data(as_text=True)).group(1)
with app.app_context():
    check("scope: the branch is stored", main.q1("SELECT branch FROM users WHERE username = 'site_a.admin'")["branch"] == "site_a")
ea = app.test_client()
login(ea, "site_a.admin", tmp_e)
ea.post("/admin/settings", data={"_csrf": token(ea, "/admin/settings"), "current": tmp_e, "new": "Site A-Admin-55", "again": "Site A-Admin-55"})
def ea_post(path, **data):
    return ea.post(path, data={"_csrf": token(ea, "/admin/staff"), **data})
page = ea.get("/admin/staff").get_data(as_text=True)
check("scope: the staff page lists only the own branch (and people without a branch)", "Site A One" in page and "Site B One" not in page)
dd9 = (date.today() + timedelta(days=900)).isoformat()
check("scope: assigning an own-branch person works", ea_post("/admin/schedule/assign", day=dd9, staff_id=e1).status_code == 302 and slot(dd9, "site_a") == e1)
check("scope: assigning a person of the other branch is refused (403)", ea_post("/admin/schedule/assign", day=dd9, staff_id=w1).status_code == 403 and slot(dd9, "site_b") is None)
post_owner("/admin/schedule/assign", day=dd9, staff_id=w1)
check("scope: clearing the other branch is refused (403)", ea_post("/admin/schedule/assign", day=dd9, staff_id="", branch="site_b").status_code == 403 and slot(dd9, "site_b") == w1)
check("scope: swap and bulk clear in the other branch are refused (403)", ea_post("/admin/schedule/swap", a=dd9, b=dd2, branch="site_b").status_code == 403 and ea_post("/admin/schedule/bulk", days=dd9, staff_id="", branch="site_b").status_code == 403)
check("scope: range for a person of the other branch is refused (403)", ea_post("/admin/schedule/range", staff_id=w1, **{"from": dd9, "to": dd9}).status_code == 403)
check("scope: switching, editing or deleting a person of the other branch is refused (403)",
      ea_post(f"/admin/staff/{w1}/branch", branch="site_a").status_code == 403 and ea_post(f"/admin/staff/{w1}/toggle").status_code == 403
      and ea_post(f"/admin/staff/{w1}/delete").status_code == 403 and ea_post("/admin/staff/save", id=str(w1), name="x", email="", phone="").status_code == 403)
ea_post("/admin/staff/save", name="New Site A Person", branch="site_b", email="", phone="")
with app.app_context():
    check("scope: a new person is always put in the administrator's own branch", main.q1("SELECT branch FROM staff WHERE name = 'New Site A Person'")["branch"] == "site_a")
page = ea.get("/admin/schedule").get_data(as_text=True)
check("scope: the other branch's lines are read only, palette lists own people only", 'class="ln w ro' in page and 'class="ln e ro' not in page and "Site B One</span>" not in page)
check("scope: the owner is not restricted", 'class="ln w ro' not in owner.get("/admin/schedule").get_data(as_text=True) and 'class="ln e ro' not in owner.get("/admin/schedule").get_data(as_text=True))
# log: an administrator sees only their own entries
page = ea.get("/admin/audit").get_data(as_text=True)
check("log: an administrator sees own entries and nobody else's", "site_a.admin" in page and 'data-label="User">owner<' not in page and 'data-label="User">nino<' not in page and 'data-label="User">admin<' not in page)
axl = _ox.load_workbook(_io.BytesIO(ea.get("/admin/audit.xlsx").data)).active
check("log: the Excel export is limited the same way", axl.max_row >= 2 and all(axl.cell(i, 2).value == "site_a.admin" for i in range(2, axl.max_row + 1)))
check("log: filtering by another user's name shows nothing of theirs", "owner" not in ea.get("/admin/audit?user=owner").get_data(as_text=True).split("<tbody>")[1])
check("log: the owner still sees everyone", "site_a.admin" in owner.get("/admin/audit").get_data(as_text=True))
# the owner changes the scope
with app.app_context():
    eid = main.q1("SELECT id FROM users WHERE username = 'site_a.admin'")["id"]
post_owner(f"/admin/users/{eid}/scope", branch="site_b")
with app.app_context():
    check("scope: the owner can change it", main.q1("SELECT branch FROM users WHERE username = 'site_a.admin'")["branch"] == "site_b")
post_owner(f"/admin/users/{eid}/scope", branch="")
with app.app_context():
    check("scope: ...and set it back to both branches", main.q1("SELECT branch FROM users WHERE username = 'site_a.admin'")["branch"] == "")
check("scope: an invalid branch is refused (400)", post_owner(f"/admin/users/{eid}/scope", branch="north").status_code == 400)
check("scope: the owner's own account cannot get a scope (400)", post_owner(f"/admin/users/{owner_id}/scope", branch="site_a").status_code == 400)
check("scope: administrators cannot change scopes (403)", ea_post(f"/admin/users/{eid}/scope", branch="site_a").status_code == 403)
check("scope: ...with both branches again the other branch is open", ea_post("/admin/schedule/assign", day=dd9, staff_id=w1).status_code == 302)

check("logout works", owner.post("/admin/logout", data={"_csrf": token(owner, "/admin/schedule")}).status_code == 302)

print()
print("ALL PASSED" if not failures else f"{len(failures)} FAILED: {failures}")
sys.exit(1 if failures else 0)
