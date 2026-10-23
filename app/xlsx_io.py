"""Excel helpers: parse a schedule pasted from Excel, and build the .xlsx export."""
import io
import re
from datetime import date

MAX_ROWS = 366

DATE_RE = re.compile(r"\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[./]\d{1,2}[./]\d{4}")
HEADERS = ("Date", "Day", "Site", "On duty", "Phone")


class TooManyRows(ValueError):
    pass


def _norm(name):
    return " ".join((name or "").split()).casefold()


def _to_date(token):
    if "-" in token:
        y, m, d = (int(x) for x in token.split("-"))
    else:
        d, m, y = (int(x) for x in re.split(r"[./]", token))
    return date(y, m, d)          # ValueError for impossible dates


def _lookup(name, by_full, by_tokens):
    """Exact full name first, then the same words in any order; None when unknown or ambiguous."""
    key = _norm(name)
    hits = by_full.get(key) or by_tokens.get(tuple(sorted(key.split())))
    return hits[0] if hits and len(hits) == 1 else None


def parse_paste(text, staff):
    """staff: iterable of {"id", "name", "branch"}. Returns {"ok": [(date, staff_id)], "bad": [(line_no, raw, reason)]}.
    Accepts 'date<TAB>name' or 'name<TAB>date'; dates as YYYY-MM-DD, DD.MM.YYYY or DD/MM/YYYY.
    The slot (Site A or Site B) is the person's branch; a person without a branch is reported, not placed."""
    lines = [(i, ln) for i, ln in enumerate((text or "").splitlines(), start=1) if ln.strip()]
    if len(lines) > MAX_ROWS:
        raise TooManyRows(len(lines))
    by_full, by_tokens, branch_of = {}, {}, {}
    for s in staff:
        key = _norm(s["name"])
        by_full.setdefault(key, []).append(s["id"])
        by_tokens.setdefault(tuple(sorted(key.split())), []).append(s["id"])
        branch_of[s["id"]] = s.get("branch", "")
    ok, bad, seen = [], [], set()
    for no, raw in lines:
        m = DATE_RE.search(raw)
        if not m:
            bad.append((no, raw.strip(), "No date found"))
            continue
        name = (raw[:m.start()] + " " + raw[m.end():]).strip(" \t;,|")
        name = " ".join(name.replace("\t", " ").split())
        if not name:
            bad.append((no, raw.strip(), "No name given"))
            continue
        try:
            day = _to_date(m.group())
        except ValueError:
            bad.append((no, raw.strip(), "Invalid date"))
            continue
        sid = _lookup(name, by_full, by_tokens)
        if sid is None:
            bad.append((no, raw.strip(), "Person not found (or the name is in the list twice)"))
            continue
        branch = branch_of.get(sid, "")
        if branch not in ("site_a", "site_b"):
            bad.append((no, raw.strip(), "This person has no site chosen (Staff page)"))
            continue
        if (day, sid) in seen:
            bad.append((no, raw.strip(), "This person is already on the list for that date"))
            continue
        seen.add((day, sid))
        ok.append((day, sid))
    return {"ok": ok, "bad": bad}


STAFF_MAX_ROWS = 200
EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
PHONE_RE = re.compile(r"[\d +()\-]{5,30}")
STAFF_HEADER_WORDS = {"full name", "name"}
BRANCH_WORDS = {"site a": "site_a", "site_a": "site_a", "a": "site_a",
                "site b": "site_b", "site_b": "site_b", "b": "site_b"}


def parse_staff_paste(text, existing):
    """Rows 'name<TAB>phone<TAB>branch' (phone and branch optional, in any order; ';' also works as separator; the branch
    is Site A or Site B). A cell with an @ is ignored, so old lists with an e-mail column still work.
    existing: iterable of {"name"}. Returns {"ok": [(name, phone, branch)], "bad": [(line_no, raw, reason)]}."""
    lines = [(i, ln) for i, ln in enumerate((text or "").splitlines(), start=1) if ln.strip()]
    if len(lines) > STAFF_MAX_ROWS:
        raise TooManyRows(len(lines))
    names = {_norm(e["name"]) for e in existing}
    ok, bad = [], []
    for pos, (no, raw) in enumerate(lines):
        parts = raw.split("\t") if "\t" in raw else raw.split(";") if ";" in raw else [raw]
        parts = [p.strip() for p in parts]
        name = " ".join(parts[0].split())
        if pos == 0 and name.casefold() in STAFF_HEADER_WORDS:
            continue
        phone, branch_text, error = "", "", ""
        for cell in parts[1:]:
            if not cell or "@" in cell:
                continue
            if cell.casefold() in BRANCH_WORDS and not branch_text:
                branch_text = cell
            elif not phone and PHONE_RE.fullmatch(cell):
                phone = cell
            elif PHONE_RE.fullmatch(cell):
                error = "The phone number is not valid"
            elif not error:
                error = "The site is not valid (Site A or Site B)" if phone or branch_text or not re.search(r"\d", cell) else "The phone number is not valid"
        if not name:
            bad.append((no, raw.strip(), "No name given"))
        elif error:
            bad.append((no, raw.strip(), error))
        elif _norm(name) in names:
            bad.append((no, raw.strip(), "This person is already in the list"))
        else:
            names.add(_norm(name))
            ok.append((name, phone, BRANCH_WORDS.get(branch_text.casefold(), "")))
    return {"ok": ok, "bad": bad}


def build_xlsx(rows, headers=HEADERS, title="Duty", widths=(14, 14, 16, 30, 18)):
    """rows: iterable of tuples matching headers -> bytes of a .xlsx file (default: the duty export)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1FAEA5")
    for r in rows:
        ws.append(list(r))
        for c in ws[ws.max_row]:
            if isinstance(c.value, str):
                c.data_type = "s"          # never let a typed "=..." become a formula
    for col, width in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
